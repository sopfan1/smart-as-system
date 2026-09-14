import streamlit as st
import pandas as pd
import sqlite3
import io
import os
import math
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage, ImageOps

# -------------------------------------------------------------
# [1. 기본 설정 및 경로 지정]
# -------------------------------------------------------------
st.set_page_config(layout="wide", page_title="Smart AS ERP - 초고속 최적화 버전")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(BASE_DIR, "attached_images")
os.makedirs(IMG_DIR, exist_ok=True)

TEMPLATE_FILE = os.path.join(BASE_DIR, "양750-03 수리 Report_260715_2.xlsx")
if not os.path.exists(TEMPLATE_FILE):
    TEMPLATE_FILE = os.path.join(BASE_DIR, "양750-03 수리 Report_260715.xlsx")

conn = sqlite3.connect("smart_as.db", check_same_thread=False)

INSPECT_1ST_PAIRS = {
    '1차_육안': '1차_육안_일자',
    '1차_특성': '1차_특성_일자',
    '1차_조합': '1차_조합_일자',
    '1차_AGING': '1차_AGING_일자',
    '1차_FULL부하': '1차_FULL부하_일자'
}

INSPECT_2ND_PAIRS = {
    '재검_육안': '재검_육안_일자',
    '재검_특성': '재검_특성_일자',
    '재검_조합': '재검_조합_일자',
    '재검_AGING': '재검_AGING_일자',
    '재검_FULL부하': '재검_FULL부하_일자'
}

EXCEL_FIELDS = [
    'NO.', '접수일', '발생일', '제조처', '접수처', '프로젝트', '제품명', '제품 S/N', '위치', '접수횟수', 
    '유/무상', '접수내역', '확인내역', 
    '1차_육안', '1차_육안_일자',
    '1차_특성', '1차_특성_일자',
    '1차_조합', '1차_조합_일자',
    '1차_AGING', '1차_AGING_일자',
    '1차_FULL부하', '1차_FULL부하_일자',
    '재검_육안', '재검_육안_일자',
    '재검_특성', '재검_특성_일자',
    '재검_조합', '재검_조합_일자',
    '재검_AGING', '재검_AGING_일자',
    '재검_FULL부하', '재검_FULL부하_일자',
    '불량원인', '수리내역', 'F/W', 'BASE PBA S/N', 'SMPS S/N', 
    'MAC', '투입부품1', '부품1 수량', '투입부품2', '부품2 수량', '투입부품3', '부품3 수량', 
    '투입부품4', '부품4 수량', '교체 (전) S/N', '교체(후) S/N', 
    '처리결과', '담당자', '완료일자', '인계일자', '비고',
    '확인내역_사진', '수리내역_사진'
]

DATE_FIELDS = [
    '접수일', '발생일', '완료일자', '인계일자',
    '1차_육안_일자', '1차_특성_일자', '1차_조합_일자', '1차_AGING_일자', '1차_FULL부하_일자',
    '재검_육안_일자', '재검_특성_일자', '재검_조합_일자', '재검_AGING_일자', '재검_FULL부하_일자'
]

# -------------------------------------------------------------
# [2. 데이터베이스 초기화 및 경량 로드 함수]
# -------------------------------------------------------------
def init_db():
    cursor = conn.cursor()
    columns_def = ", ".join([f'"{col}" TEXT' for col in EXCEL_FIELDS])
    cursor.execute(f'''
        CREATE TABLE IF NOT EXISTS as_data (
            {columns_def}
        )
    ''')
    conn.commit()

init_db()

def load_fresh_db_data(order_mode="최신순 (마지막 NO.부터)"):
    try:
        df = pd.read_sql("SELECT rowid as rowid_val, * FROM as_data", conn)
    except:
        init_db()
        df = pd.read_sql("SELECT rowid as rowid_val, * FROM as_data", conn)
        
    for col in EXCEL_FIELDS:
        if col not in df.columns: df[col] = ""
    ed_df = df[['rowid_val'] + EXCEL_FIELDS].copy()
    
    for col in DATE_FIELDS:
        if col in ed_df.columns:
            ed_df[col] = ed_df[col].astype(str).str.strip().replace(['nan', 'None', 'NAT', 'NaT'], '')
    
    for col in list(INSPECT_1ST_PAIRS.keys()) + list(INSPECT_2ND_PAIRS.keys()):
        ed_df[col] = ed_df[col].fillna("").astype(str).str.strip().str.upper()
        ed_df[col] = ed_df[col].apply(lambda x: x if x in ['PASS', 'FAIL'] else "")
        
    ed_df['temp_no_sort'] = ed_df['NO.'].apply(lambda x: safe_int_no(x) if safe_int_no(x) is not None else 999999)
    ed_df.sort_values(by='temp_no_sort', ascending=True, inplace=True)
    ed_df.drop(columns=['temp_no_sort'], inplace=True)

    if order_mode.startswith("최신순"):
        ed_df = ed_df.iloc[::-1].reset_index(drop=True)
    else:
        ed_df = ed_df.reset_index(drop=True)
    return ed_df

# -------------------------------------------------------------
# [3. 보조 유틸 함수]
# -------------------------------------------------------------
def clean_date_str(val):
    if pd.isna(val): return ""
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ['nan', 'none', 'nat', 'null']: return ""
    if " " in val_str and "\n" not in val_str and "~" not in val_str: val_str = val_str.split(" ")[0]
    if "T" in val_str and "\n" not in val_str and "~" not in val_str: val_str = val_str.split("T")[0]
    return val_str

def calc_aging_48h(start_date_str, multiline=False):
    if not start_date_str: return ""
    raw = str(start_date_str).strip()
    if "~" in raw: return raw
    try:
        dt = datetime.strptime(raw[:10], "%Y-%m-%d")
        dt_end = dt + timedelta(days=2)
        s_str = dt.strftime("%Y-%m-%d")
        e_str = dt_end.strftime("%Y-%m-%d")
        if multiline:
            return f"{s_str}\n~\n{e_str}"
        return f"{s_str} ~ {e_str}"
    except:
        return raw

def safe_int_no(val):
    if pd.isna(val): return None
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ['nan', 'none', 'null', '']: return None
    if "(" in val_str: val_str = val_str.split("(")[0].strip()
    try:
        float_val = float(val_str)
        if math.isnan(float_val): return None
        return int(float_val)
    except:
        return None

def calculate_reception_counts(df_to_calc):
    sn_counter = {}
    counts = []
    seen_no_map = {}
    
    for _, row in df_to_calc.iterrows():
        sn_clean = str(row.get('제품 S/N', '')).strip()
        no_val = str(row.get('NO.', '')).strip()
        
        if not sn_clean or sn_clean.lower() in ['nan', 'none', '-', '']:
            counts.append("1")
        else:
            if no_val and no_val in seen_no_map:
                counts.append(seen_no_map[no_val])
            else:
                sn_counter[sn_clean] = sn_counter.get(sn_clean, 0) + 1
                assigned_count = str(sn_counter[sn_clean])
                if no_val: seen_no_map[no_val] = assigned_count
                counts.append(assigned_count)
                
    df_to_calc['접수횟수'] = counts
    return df_to_calc

def make_excel_bytes_with_merge(dataframe):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'AS관리대장'

    excel_headers = [
        'NO.', '접수일', '발생일', '제조처', '접수처', '프로젝트', '제품명', '제품 S/N', '위치', '접수횟수', 
        '유/무상', '접수내역', '확인내역', 
        '검사_육안', '검사_육안_일자', '검사_전자소자', '검사_전자소자_일자', 
        '검사_기능조합', '검사_기능조합_일자', '검사_AGING', '검사_AGING_일자', 
        '검사_FULL부하', '검사_FULL부하_일자',
        '불량원인', '수리내역', 'F/W', 'BASE PBA S/N', 'SMPS S/N', 
        'MAC', '투입부품1', '부품1 수량', '투입부품2', '부품2 수량', '투입부품3', '부품3 수량', 
        '투입부품4', '부품4 수량', '교체 (전) S/N', '교체(후) S/N', 
        '처리결과', '담당자', '완료일자', '인계일자', '비고'
    ]
    ws.append(excel_headers)

    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'), right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'), bottom=Side(style='thin', color='D3D3D3')
    )
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    for col_num in range(1, len(excel_headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = openpyxl.styles.Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    curr_row = 2
    for _, r in dataframe.iterrows():
        r_dict = r.to_dict()
        has_retest = any(str(r_dict.get(k, '')).strip() != '' for k in list(INSPECT_2ND_PAIRS.keys()) + list(INSPECT_2ND_PAIRS.values()))

        row1_vals = [
            r_dict.get('NO.', ''), r_dict.get('접수일', ''), r_dict.get('발생일', ''), r_dict.get('제조처', ''),
            r_dict.get('접수처', ''), r_dict.get('프로젝트', ''), r_dict.get('제품명', ''), r_dict.get('제품 S/N', ''),
            r_dict.get('위치', ''), r_dict.get('접수횟수', ''), r_dict.get('유/무상', ''), r_dict.get('접수내역', ''),
            r_dict.get('확인내역', ''),
            r_dict.get('1차_육안', ''), r_dict.get('1차_육안_일자', ''),
            r_dict.get('1차_특성', ''), r_dict.get('1차_특성_일자', ''),
            r_dict.get('1차_조합', ''), r_dict.get('1차_조합_일자', ''),
            r_dict.get('1차_AGING', ''), calc_aging_48h(r_dict.get('1차_AGING_일자', '')),
            r_dict.get('1차_FULL부하', ''), r_dict.get('1차_FULL부하_일자', ''),
            r_dict.get('불량원인', ''), r_dict.get('수리내역', ''), r_dict.get('F/W', ''),
            r_dict.get('BASE PBA S/N', ''), r_dict.get('SMPS S/N', ''), r_dict.get('MAC', ''),
            r_dict.get('투입부품1', ''), r_dict.get('부품1 수량', ''), r_dict.get('투입부품2', ''), r_dict.get('부품2 수량', ''),
            r_dict.get('투입부품3', ''), r_dict.get('부품3 수량', ''), r_dict.get('투입부품4', ''), r_dict.get('부품4 수량', ''),
            r_dict.get('교체 (전) S/N', ''), r_dict.get('교체(후) S/N', ''),
            r_dict.get('처리결과', ''), r_dict.get('담당자', ''), r_dict.get('완료일자', ''), r_dict.get('인계일자', ''),
            r_dict.get('비고', '')
        ]

        if has_retest:
            row2_vals = list(row1_vals)
            row2_vals[13] = r_dict.get('재검_육안', '')
            row2_vals[14] = r_dict.get('재검_육안_일자', '')
            row2_vals[15] = r_dict.get('재검_특성', '')
            row2_vals[16] = r_dict.get('재검_특성_일자', '')
            row2_vals[17] = r_dict.get('재검_조합', '')
            row2_vals[18] = r_dict.get('재검_조합_일자', '')
            row2_vals[19] = r_dict.get('재검_AGING', '')
            row2_vals[20] = calc_aging_48h(r_dict.get('재검_AGING_일자', ''))
            row2_vals[21] = r_dict.get('재검_FULL부하', '')
            row2_vals[22] = r_dict.get('재검_FULL부하_일자', '')

            for col_i, v in enumerate(row1_vals, 1):
                c = ws.cell(row=curr_row, column=col_i, value=str(v))
                c.border = thin_border
                c.alignment = Alignment(vertical="center", horizontal="center" if col_i in [1, 10] or "일자" in excel_headers[col_i-1] or (14 <= col_i <= 23) else "left")

            for col_i, v in enumerate(row2_vals, 1):
                c = ws.cell(row=curr_row + 1, column=col_i, value=str(v))
                c.border = thin_border
                if 14 <= col_i <= 23:
                    c.fill = yellow_fill
                c.alignment = Alignment(vertical="center", horizontal="center" if col_i in [1, 10] or "일자" in excel_headers[col_i-1] or (14 <= col_i <= 23) else "left")

            for col_i in list(range(1, 14)) + list(range(24, len(excel_headers) + 1)):
                ws.merge_cells(start_row=curr_row, end_row=curr_row + 1, start_column=col_i, end_column=col_i)

            curr_row += 2
        else:
            for col_i, v in enumerate(row1_vals, 1):
                c = ws.cell(row=curr_row, column=col_i, value=str(v))
                c.border = thin_border
                c.alignment = Alignment(vertical="center", horizontal="center" if col_i in [1, 10] or "일자" in excel_headers[col_i-1] or (14 <= col_i <= 23) else "left")
            curr_row += 1

    out_buf = io.BytesIO()
    wb.save(out_buf)
    return out_buf.getvalue()

def prepare_excel_image(filepath, target_w_cm=4.5, target_h_cm=6.0):
    if not filepath or pd.isna(filepath): return None
    fp = str(filepath).strip()
    if not fp or fp.lower() in ['nan', 'none', '']: return None

    resolved_path = None
    if os.path.exists(fp):
        resolved_path = fp
    else:
        base_name = os.path.basename(fp)
        p1 = os.path.join(IMG_DIR, base_name)
        p2 = os.path.join(BASE_DIR, base_name)
        if os.path.exists(p1): resolved_path = p1
        elif os.path.exists(p2): resolved_path = p2
        else:
            if os.path.exists(IMG_DIR):
                for img_f in os.listdir(IMG_DIR):
                    if base_name in img_f or img_f in base_name:
                        resolved_path = os.path.join(IMG_DIR, img_f)
                        break

    if not resolved_path or not os.path.exists(resolved_path):
        return None

    try:
        target_w = round((target_w_cm / 2.54) * 96)
        target_h = round((target_h_cm / 2.54) * 96)
        with PILImage.open(resolved_path) as raw_img:
            img = ImageOps.exif_transpose(raw_img)
            img_resized = img.resize((target_w, target_h), PILImage.Resampling.LANCZOS)
            buf = io.BytesIO()
            img_resized.convert("RGB").save(buf, format='JPEG', quality=95)
            buf.seek(0)
            xl_img = OpenpyxlImage(buf)
            xl_img.width = target_w
            xl_img.height = target_h
            xl_img._image_buffer = buf
            return xl_img
    except: return None

def set_safe_cell_value(ws, cell, value):
    target_cell = cell
    if type(cell).__name__ == 'MergedCell':
        for m in ws.merged_cells.ranges:
            if cell.coordinate in m:
                target_cell = ws.cell(row=m.min_row, column=m.min_col)
                break
    target_cell.value = str(value) if value is not None else ""
    if "\n" in str(value):
        target_cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
    else:
        target_cell.alignment = Alignment(vertical="center", horizontal="left")

def get_adjacent_input_cell(ws, label_cell, offset_col=1):
    target_col = label_cell.column + offset_col
    target_row = label_cell.row
    for m in ws.merged_cells.ranges:
        if label_cell.coordinate in m:
            target_col = m.max_col + offset_col
            target_row = m.min_row
            break
    return ws.cell(row=target_row, column=target_col)

def add_image_with_nudge(ws, xl_img, top_col_1based, top_row_1based, nudge_x=4, nudge_y=4):
    col_idx = top_col_1based - 1
    row_idx = top_row_1based - 1
    offset_x_emu = nudge_x * 12700
    offset_y_emu = nudge_y * 12700
    size = XDRPositiveSize2D(cx=pixels_to_EMU(xl_img.width), cy=pixels_to_EMU(xl_img.height))
    marker = AnchorMarker(col=col_idx, colOff=offset_x_emu, row=row_idx, rowOff=offset_y_emu)
    xl_img.anchor = OneCellAnchor(_from=marker, ext=size)
    ws.add_image(xl_img)

def format_report_result(val):
    if not val: return "-"
    lines = [v.strip().upper() for v in str(val).split("\n") if v.strip()]
    formatted = []
    for l in lines:
        if "PASS" in l: formatted.append("PASS")
        elif "FAIL" in l or "NG" in l: formatted.append("NG")
        else: formatted.append("-")
    return "\n".join(formatted) if formatted else "-"

def generate_multi_repair_excel(selected_rows_data, template_filename=TEMPLATE_FILE):
    if not os.path.exists(template_filename):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Report"
        out_buf = io.BytesIO()
        wb.save(out_buf)
        return out_buf.getvalue()

    wb = openpyxl.load_workbook(template_filename)
    template_ws = wb.active

    active_buffers = []

    for idx, row_data in enumerate(selected_rows_data):
        no_val = str(row_data.get('NO.', idx + 1)).replace('/', '_').strip()
        if not no_val: no_val = str(idx + 1)
        sheet_title = f"Report_NO_{no_val}"[:30]

        if idx == 0:
            ws = template_ws
            ws.title = sheet_title
        else:
            ws = wb.copy_worksheet(template_ws)
            ws.title = sheet_title

        photo_cells_map = {}

        def merge_lines(val1, val2):
            v1 = str(val1).strip()
            v2 = str(val2).strip()
            if v1 and v2: return f"{v1}\n{v2}"
            return v1 or v2 or ""

        aging_1st_str = calc_aging_48h(row_data.get('1차_AGING_일자', ''), multiline=True)
        aging_2nd_str = calc_aging_48h(row_data.get('재검_AGING_일자', ''), multiline=True)

        inspect_merged = {
            '육안_결과': merge_lines(row_data.get('1차_육안', ''), row_data.get('재검_육안', '')),
            '육안_일자': merge_lines(row_data.get('1차_육안_일자', ''), row_data.get('재검_육안_일자', '')),
            '특성_결과': merge_lines(row_data.get('1차_특성', ''), row_data.get('재검_특성', '')),
            '특성_일자': merge_lines(row_data.get('1차_특성_일자', ''), row_data.get('재검_특성_일자', '')),
            '조합_결과': merge_lines(row_data.get('1차_조합', ''), row_data.get('재검_조합', '')),
            '조합_일자': merge_lines(row_data.get('1차_조합_일자', ''), row_data.get('재검_조합_일자', '')),
            'AGING_결과': merge_lines(row_data.get('1차_AGING', ''), row_data.get('재검_AGING', '')),
            'AGING_일자': merge_lines(aging_1st_str, aging_2nd_str),
            'FULL부하_결과': merge_lines(row_data.get('1차_FULL부하', ''), row_data.get('재검_FULL부하', '')),
            'FULL부하_일자': merge_lines(row_data.get('1차_FULL부하_일자', ''), row_data.get('재검_FULL부하_일자', ''))
        }

        defect_cause = str(row_data.get('불량원인', '')).strip()
        repair_action = str(row_data.get('수리내역', '')).strip()
        
        if defect_cause and repair_action:
            combined_repair_desc = f"[불량원인] {defect_cause}\n[수리내역] {repair_action}"
        elif defect_cause:
            combined_repair_desc = f"[불량원인] {defect_cause}"
        else:
            combined_repair_desc = repair_action

        date_header_col = None
        result_header_col = None

        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                val_check = str(ws.cell(row=r, column=c).value or '').replace(' ', '')
                if '검사일자' in val_check:
                    date_header_col = c
                elif val_check == '판정':
                    result_header_col = c

        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                val = str(cell.value).strip() if cell.value is not None else ""
                val_no_space = val.replace(" ", "")

                if not val: continue

                if val == "접수일":
                    inp_cell = get_adjacent_input_cell(ws, cell)
                    set_safe_cell_value(ws, inp_cell, clean_date_str(row_data.get('접수일', '')))
                elif val == "프로젝트":
                    inp_cell = get_adjacent_input_cell(ws, cell)
                    set_safe_cell_value(ws, inp_cell, row_data.get('프로젝트', ''))
                elif val == "제품명":
                    inp_cell = get_adjacent_input_cell(ws, cell)
                    set_safe_cell_value(ws, inp_cell, row_data.get('제품명', ''))
                elif "S/N" in val.upper() or "제품S/N" in val:
                    inp_cell = get_adjacent_input_cell(ws, cell)
                    set_safe_cell_value(ws, inp_cell, row_data.get('제품 S/N', ''))
                elif val == "접수내역":
                    inp_cell = get_adjacent_input_cell(ws, cell)
                    set_safe_cell_value(ws, inp_cell, row_data.get('접수내역', ''))
                elif "불량증상기재" in val_no_space or "불량증상" in val_no_space or "불량증상기재" in val:
                    target_row = cell.row
                    j_cell = ws.cell(row=target_row, column=10)
                    set_safe_cell_value(ws, j_cell, row_data.get('확인내역', ''))
                elif val == "수리내역" or "수리내역기재" in val_no_space:
                    target_row = cell.row
                    j_cell = ws.cell(row=target_row, column=10)
                    set_safe_cell_value(ws, j_cell, combined_repair_desc)
                elif val == "비고":
                    inp_cell = get_adjacent_input_cell(ws, cell)
                    set_safe_cell_value(ws, inp_cell, row_data.get('비고', ''))
                elif "확인내역_사진" in val_no_space:
                    photo_cells_map['conf'] = cell
                elif "수리내역_사진" in val_no_space:
                    photo_cells_map['rep'] = cell

                elif "육안검사" in val:
                    d_col = date_header_col if date_header_col else (cell.column + 2)
                    res_col = result_header_col if result_header_col else (d_col + 1)
                    set_safe_cell_value(ws, ws.cell(row=r, column=d_col), inspect_merged['육안_일자'] or "-")
                    set_safe_cell_value(ws, ws.cell(row=r, column=res_col), format_report_result(inspect_merged['육안_결과']))

                elif "전자소자" in val or "특성검사" in val:
                    d_col = date_header_col if date_header_col else (cell.column + 2)
                    res_col = result_header_col if result_header_col else (d_col + 1)
                    set_safe_cell_value(ws, ws.cell(row=r, column=d_col), inspect_merged['특성_일자'] or "-")
                    set_safe_cell_value(ws, ws.cell(row=r, column=res_col), format_report_result(inspect_merged['특성_결과']))

                elif "기능/조합" in val or "조합시험" in val:
                    d_col = date_header_col if date_header_col else (cell.column + 2)
                    res_col = result_header_col if result_header_col else (d_col + 1)
                    set_safe_cell_value(ws, ws.cell(row=r, column=d_col), inspect_merged['조합_일자'] or "-")
                    set_safe_cell_value(ws, ws.cell(row=r, column=res_col), format_report_result(inspect_merged['조합_결과']))

                elif "AGING" in val.upper() or "에이징" in val:
                    d_col = date_header_col if date_header_col else (cell.column + 2)
                    res_col = result_header_col if result_header_col else (d_col + 1)
                    set_safe_cell_value(ws, ws.cell(row=r, column=d_col), inspect_merged['AGING_일자'] or "-")
                    set_safe_cell_value(ws, ws.cell(row=r, column=res_col), format_report_result(inspect_merged['AGING_결과']))

                elif "FULL부하" in val.upper() or "LONGRUN" in val.upper():
                    d_col = date_header_col if date_header_col else (cell.column + 2)
                    res_col = result_header_col if result_header_col else (d_col + 1)
                    set_safe_cell_value(ws, ws.cell(row=r, column=d_col), inspect_merged['FULL부하_일자'] or "-")
                    set_safe_cell_value(ws, ws.cell(row=r, column=res_col), format_report_result(inspect_merged['FULL부하_결과']))

        if 'conf' in photo_cells_map:
            p_cell1 = photo_cells_map['conf']
            top_r, top_c = p_cell1.row, p_cell1.column
            for m in ws.merged_cells.ranges:
                if p_cell1.coordinate in m:
                    top_r, top_c = m.min_row, m.min_col
                    break
            ws.cell(row=top_r, column=top_c).value = ""
            img1 = prepare_excel_image(row_data.get('확인내역_사진'))
            if img1:
                if hasattr(img1, '_image_buffer'): active_buffers.append(img1._image_buffer)
                add_image_with_nudge(ws, img1, top_c, top_r, nudge_x=4, nudge_y=4)

        if 'rep' in photo_cells_map:
            p_cell2 = photo_cells_map['rep']
            top_r, top_c = p_cell2.row, p_cell2.column
            for m in ws.merged_cells.ranges:
                if p_cell2.coordinate in m:
                    top_r, top_c = m.min_row, m.min_col
                    break
            ws.cell(row=top_r, column=top_c).value = ""
            img2 = prepare_excel_image(row_data.get('수리내역_사진'))
            if img2:
                if hasattr(img2, '_image_buffer'): active_buffers.append(img2._image_buffer)
                add_image_with_nudge(ws, img2, top_c, top_r, nudge_x=4, nudge_y=4)

    out_buf = io.BytesIO()
    wb.save(out_buf)
    return out_buf.getvalue()

# -------------------------------------------------------------
# [4. 사이드바 로그인 및 권한 관리 시스템]
# -------------------------------------------------------------
st.sidebar.title("🔐 사용자 인증 및 권한")
auth_mode = st.sidebar.radio("접속 모드 선택", ["일반 사용자 (조회/REPORT 출력)", "마스터 관리자 (등록/수정 권한)"])

is_master = False
if auth_mode == "마스터 관리자 (등록/수정 권한)":
    input_id = st.sidebar.text_input("관리자 ID 입력", value="")
    input_pw = st.sidebar.text_input("관리자 비밀번호 입력", type="password", value="")
    
    if input_id == "jwko" and input_pw == "qcteam12!":
        is_master = True
        st.sidebar.success("✅ 마스터 관리자 인증 완료 (수정/등록 가능)")
    else:
        if input_id or input_pw:
            st.sidebar.error("❌ ID 또는 비밀번호가 일치하지 않습니다.")
        else:
            st.sidebar.warning("⚠️ 관리자 ID와 비밀번호를 입력하세요.")
        is_master = False
else:
    st.sidebar.info("👁️ 일반 사용자 모드: 대장 조회 및 수리 REPORT 발행만 가능합니다.")

st.title("🏢 Smart AS Management & KPI System")

if "last_order" not in st.session_state:
    st.session_state["last_order"] = "최신순 (마지막 NO.부터)"

current_order = st.radio("대장 표시 순서", ["최신순 (마지막 NO.부터)", "과거순 (NO. 1부터)"], horizontal=True, key="sort_order_radio")

if current_order != st.session_state["last_order"]:
    st.session_state["last_order"] = current_order
    st.session_state["display_df"] = load_fresh_db_data(current_order)
    if '선택' not in st.session_state["display_df"].columns:
        st.session_state["display_df"].insert(0, '선택', False)

if "display_df" not in st.session_state or st.session_state["display_df"] is None:
    st.session_state["display_df"] = load_fresh_db_data(current_order)
    if '선택' not in st.session_state["display_df"].columns:
        st.session_state["display_df"].insert(0, '선택', False)

if is_master:
    with st.expander("📥 [관리자 전용] 엑셀 파일 업로드 및 DB 동기화", expanded=False):
        uploaded_file = st.file_uploader("AS관리대장 엑셀 파일(.xlsx) 선택", type=["xlsx", "xls"], key="excel_uploader")
        if uploaded_file is not None and st.button("🚀 엑셀 데이터 DB로 자동 기입"):
            try:
                excel_df = pd.read_excel(uploaded_file)
                excel_df.columns = [str(c).strip() for c in excel_df.columns]
                rename_dict = {}
                for col in excel_df.columns:
                    clean_name = col.replace(" ", "").upper()
                    if clean_name in ['NO.', 'NO', '순번', '번호', 'INDEX']:
                        rename_dict[col] = 'NO.'
                        break
                if rename_dict: excel_df.rename(columns=rename_dict, inplace=True)
                for col in EXCEL_FIELDS:
                    if col not in excel_df.columns: excel_df[col] = ""
                
                excel_df['NO.'] = [str(safe_int_no(v)) if safe_int_no(v) is not None else str(i+1) for i, v in enumerate(excel_df['NO.'])]
                for col in DATE_FIELDS:
                    if col in excel_df.columns: excel_df[col] = excel_df[col].apply(clean_date_str)
                excel_df = calculate_reception_counts(excel_df)
                final_import_df = excel_df[EXCEL_FIELDS].fillna("").astype(str)
                final_import_df.to_sql("as_data", conn, if_exists="replace", index=False)
                
                st.session_state["display_df"] = load_fresh_db_data(current_order)
                if '선택' not in st.session_state["display_df"].columns:
                    st.session_state["display_df"].insert(0, '선택', False)
                st.success("데이터베이스에 정상 반영되었습니다.")
                st.rerun()
            except Exception as ex:
                st.error(f"오류: {ex}")

        st.markdown("---")
        if st.button("⚠️ 모든 관리대장 데이터 전체 초기화 (처음부터 다시 시작)", type="secondary"):
            try:
                conn.execute("DELETE FROM as_data")
                conn.commit()
                
                st.session_state["display_df"] = load_fresh_db_data(current_order)
                if '선택' not in st.session_state["display_df"].columns:
                    st.session_state["display_df"].insert(0, '선택', False)
                    
                st.toast("🗑️ 모든 데이터가 성공적으로 초기화되었습니다.", icon="✅")
                st.rerun()
            except Exception as reset_err:
                st.error(f"초기화 중 오류 발생: {reset_err}")

try:
    tab1, tab2 = st.tabs(["📝 AS 관리대장 & 수리 REPORT 발행", "📊 2026년 KPI 데이터 분석"])

    # =============================================================
    # [TAB 1: 관리대장 및 수리 REPORT 발행]
    # =============================================================
    with tab1:
        if is_master:
            with st.expander("➕ [관리자 전용] 신규 AS 접수 데이터 즉시 등록", expanded=False):
                with st.form("new_as_entry_form", clear_on_submit=True):
                    st.markdown("##### 📝 신규 접수 기본 정보 입력")
                    n_c1, n_c2, n_c3, n_c4 = st.columns(4)
                    with n_c1:
                        new_date = st.text_input("접수일", value=datetime.now().strftime('%Y-%m-%d'))
                        new_project = st.text_input("프로젝트", placeholder="예: SCAM 모듈")
                    with n_c2:
                        new_prod_name = st.text_input("제품명", placeholder="예: CPU PBA")
                        new_sn = st.text_input("제품 S/N", placeholder="제품 일련번호")
                    with n_c3:
                        new_client = st.text_input("접수처", placeholder="고객사/협력사")
                        new_cost = st.selectbox("유/무상", ["무상", "유상", "기타"], index=0)
                    with n_c4:
                        new_manager = st.text_input("담당자", placeholder="처리 담당자명")
                        new_res = st.text_input("처리결과", value="접수")

                    st.markdown("##### 🔍 상세 증상 및 수리 조치")
                    nd_c1, nd_c2 = st.columns(2)
                    with nd_c1:
                        new_recept_desc = st.text_area("접수내역", height=70)
                        new_defect_cause = st.text_area("불량원인", height=70)
                    with nd_c2:
                        new_check_desc = st.text_area("확인내역", height=70)
                        new_repair_desc = st.text_area("수리내역", height=70)

                    submit_btn = st.form_submit_button("🚀 신규 데이터 등록")

                    if submit_btn:
                        if not new_sn.strip() and not new_prod_name.strip():
                            st.warning("제품명 또는 제품 S/N을 최소 1개 이상 입력해 주세요.")
                        else:
                            current_all_df = load_fresh_db_data(current_order)
                            int_nos = [safe_int_no(v) for v in current_all_df['NO.']]
                            valid_nos = [x for x in int_nos if x is not None]
                            next_no = (max(valid_nos) + 1) if valid_nos else 1

                            new_entry = {col: "" for col in EXCEL_FIELDS}
                            new_entry['NO.'] = str(next_no)
                            new_entry['접수일'] = clean_date_str(new_date)
                            new_entry['프로젝트'] = str(new_project).strip()
                            new_entry['제품명'] = str(new_prod_name).strip()
                            new_entry['제품 S/N'] = str(new_sn).strip()
                            new_entry['접수처'] = str(new_client).strip()
                            new_entry['유/무상'] = str(new_cost).strip()
                            new_entry['담당자'] = str(new_manager).strip()
                            new_entry['처리결과'] = str(new_res).strip()
                            new_entry['접수내역'] = str(new_recept_desc).strip()
                            new_entry['확인내역'] = str(new_check_desc).strip()
                            new_entry['불량원인'] = str(new_defect_cause).strip()
                            new_entry['수리내역'] = str(new_repair_desc).strip()

                            new_row_df = pd.DataFrame([new_entry])
                            combined_df = pd.concat([current_all_df[EXCEL_FIELDS], new_row_df], ignore_index=True)
                            combined_df = calculate_reception_counts(combined_df)

                            combined_df[EXCEL_FIELDS].fillna("").astype(str).to_sql("as_data", conn, if_exists="replace", index=False)
                            
                            st.session_state["display_df"] = load_fresh_db_data(current_order)
                            if '선택' not in st.session_state["display_df"].columns:
                                st.session_state["display_df"].insert(0, '선택', False)
                            st.toast(f"✅ NO.{next_no} 등록 완료!", icon="🎉")
                            st.rerun()

            with st.expander("📷 [관리자 전용] 특정 접수 건 사진 등록 및 삭제", expanded=False):
                fresh_check = load_fresh_db_data(current_order)
                if not fresh_check.empty:
                    valid_photo_nos = [str(v) for v in fresh_check['NO.'].tolist() if str(v).strip()]
                    if valid_photo_nos:
                        sel_no = st.selectbox("사진 관리할 접수 NO. 선택:", valid_photo_nos, key="sel_no_box")
                        target_matches = fresh_check[fresh_check['NO.'] == sel_no]
                        if not target_matches.empty:
                            target_idx = target_matches.index[0]
                            cur_p1 = fresh_check.loc[target_idx, '확인내역_사진']
                            cur_p2 = fresh_check.loc[target_idx, '수리내역_사진']
                            
                            st.caption(f"📌 **NO.{sel_no} 현재 등록 상태** — 불량증상 사진: {'[등록됨]' if cur_p1 else '[없음]'} | 수리내역 사진: {'[등록됨]' if cur_p2 else '[없음]'}")

                            c1, c2 = st.columns(2)
                            with c1:
                                up_img1 = st.file_uploader(f"NO.{sel_no} 불량 증상 사진 업로드", type=["png", "jpg", "jpeg"], key=f"p1_{sel_no}")
                            with c2:
                                up_img2 = st.file_uploader(f"NO.{sel_no} 수리 내역 사진 업로드", type=["png", "jpg", "jpeg"], key=f"p2_{sel_no}")

                            sc1, sc2 = st.columns(2)
                            with sc1:
                                if st.button("💾 사진 DB 저장 / 갱신"):
                                    if up_img1:
                                        path1 = os.path.join(IMG_DIR, f"defect_{sel_no}_{datetime.now().strftime('%H%M%S')}.jpg")
                                        with open(path1, "wb") as f: f.write(up_img1.getbuffer())
                                        fresh_check.at[target_idx, '확인내역_사진'] = path1
                                    if up_img2:
                                        path2 = os.path.join(IMG_DIR, f"repair_{sel_no}_{datetime.now().strftime('%H%M%S')}.jpg")
                                        with open(path2, "wb") as f: f.write(up_img2.getbuffer())
                                        fresh_check.at[target_idx, '수리내역_사진'] = path2
                                    
                                    fresh_check[EXCEL_FIELDS].to_sql("as_data", conn, if_exists="replace", index=False)
                                    st.session_state["display_df"] = load_fresh_db_data(current_order)
                                    if '선택' not in st.session_state["display_df"].columns:
                                        st.session_state["display_df"].insert(0, '선택', False)
                                    st.toast("사진이 정상 저장되었습니다!", icon="✅")
                                    st.rerun()

                            with sc2:
                                if st.button("🗑️ 등록된 사진 전체 삭제"):
                                    fresh_check.at[target_idx, '확인내역_사진'] = ""
                                    fresh_check.at[target_idx, '수리내역_사진'] = ""
                                    fresh_check[EXCEL_FIELDS].to_sql("as_data", conn, if_exists="replace", index=False)
                                    st.session_state["display_df"] = load_fresh_db_data(current_order)
                                    if '선택' not in st.session_state["display_df"].columns:
                                        st.session_state["display_df"].insert(0, '선택', False)
                                    st.toast("해당 건의 사진이 모두 삭제되었습니다.", icon="🗑️")
                                    st.rerun()
        else:
            st.info("ℹ️ 현재 **일반 사용자 모드**입니다. 신규 등록 및 데이터 수정은 사이드바에서 마스터 계정(`jwko`)으로 로그인 후 가능합니다.")

        # -------------------------------------------------------------
        # [엑셀 스타일 검색 및 필터 바]
        # -------------------------------------------------------------
        st.markdown("#### 🔎 AS관리대장 검색 및 필터 바")
        s_c1, s_c2, s_c3, s_c4 = st.columns(4)
        with s_c1:
            search_query = st.text_input("🔍 통합 검색 (제품명, S/N 등)", value="", key="ledger_search_box")
        with s_c2:
            all_projs = sorted([str(p).strip() for p in st.session_state["display_df"]['프로젝트'].unique() if str(p).strip()])
            sel_proj_filter = st.selectbox("🏗️ 프로젝트 필터", ["전체 프로젝트"] + all_projs, key="ledger_proj_filter")
        with s_c3:
            sel_cost_filter = st.selectbox("💰 유/무상 필터", ["전체", "무상", "유상"], key="ledger_cost_filter")
        with s_c4:
            all_res = sorted([str(r).strip() for r in st.session_state["display_df"]['처리결과'].unique() if str(r).strip()])
            sel_res_filter = st.selectbox("📌 처리결과 필터", ["전체"] + all_res, key="ledger_res_filter")

        view_df = st.session_state["display_df"].copy()
        if search_query.strip():
            q = search_query.strip().lower()
            mask = view_df.astype(str).apply(lambda row: row.str.lower().str.contains(q).any(), axis=1)
            view_df = view_df[mask]
        if sel_proj_filter != "전체 프로젝트":
            view_df = view_df[view_df['프로젝트'].astype(str).str.strip() == sel_proj_filter]
        if sel_cost_filter != "전체":
            view_df = view_df[view_df['유/무상'].astype(str).str.strip().str.contains(sel_cost_filter)]
        if sel_res_filter != "전체":
            view_df = view_df[view_df['처리결과'].astype(str).str.strip() == sel_res_filter]

        st.caption(f"총 데이터: {len(st.session_state['display_df']):,}건 중 **검색/필터된 항목: {len(view_df):,}건** 표시 중")

        if is_master:
            st.markdown("#### 📋 AS관리대장 (관리자 편집 모드)")
        else:
            st.markdown("#### 📋 AS관리대장 (조회 전용 모드)")

        pass_fail_options = ["", "PASS", "FAIL"]
        
        column_config = {
            "선택": st.column_config.CheckboxColumn("선택", help="수리 REPORT 발행 시 체크", default=False),
            "NO.": st.column_config.TextColumn("NO.", disabled=not is_master),
            "접수횟수": st.column_config.TextColumn("접수횟수", disabled=True),
            "1차_육안": st.column_config.SelectboxColumn("1차_육안", options=pass_fail_options, required=False, disabled=not is_master),
            "1차_육안_일자": st.column_config.TextColumn("1차_육안_일자", disabled=not is_master),
            "1차_특성": st.column_config.SelectboxColumn("1차_특성", options=pass_fail_options, required=False, disabled=not is_master),
            "1차_특성_일자": st.column_config.TextColumn("1차_특성_일자", disabled=not is_master),
            "1차_조합": st.column_config.SelectboxColumn("1차_조합", options=pass_fail_options, required=False, disabled=not is_master),
            "1차_조합_일자": st.column_config.TextColumn("1차_조합_일자", disabled=not is_master),
            "1차_AGING": st.column_config.SelectboxColumn("1차_AGING", options=pass_fail_options, required=False, disabled=not is_master),
            "1차_AGING_일자": st.column_config.TextColumn("1차_AGING_일자", disabled=not is_master),
            "1차_FULL부하": st.column_config.SelectboxColumn("1차_FULL부하", options=pass_fail_options, required=False, disabled=not is_master),
            "1차_FULL부하_일자": st.column_config.TextColumn("1차_FULL부하_일자", disabled=not is_master),
            "재검_육안": st.column_config.SelectboxColumn("🟡 재검_육안", options=pass_fail_options, required=False, disabled=not is_master),
            "재검_육안_일자": st.column_config.TextColumn("🟡 재검_육안_일자", disabled=not is_master),
            "재검_특성": st.column_config.SelectboxColumn("🟡 재검_특성", options=pass_fail_options, required=False, disabled=not is_master),
            "재검_특성_일자": st.column_config.TextColumn("🟡 재검_특성_일자", disabled=not is_master),
            "재검_조합": st.column_config.SelectboxColumn("🟡 재검_조합", options=pass_fail_options, required=False, disabled=not is_master),
            "재검_조합_일자": st.column_config.TextColumn("🟡 재검_조합_일자", disabled=not is_master),
            "재검_AGING": st.column_config.SelectboxColumn("🟡 재검_AGING", options=pass_fail_options, required=False, disabled=not is_master),
            "재검_AGING_일자": st.column_config.TextColumn("🟡 재검_AGING_일자", disabled=not is_master),
            "재검_FULL부하": st.column_config.SelectboxColumn("🟡 재검_FULL부하", options=pass_fail_options, required=False, disabled=not is_master),
            "재검_FULL부하_일자": st.column_config.TextColumn("🟡 재검_FULL부하_일자", disabled=not is_master),
            "확인내역": st.column_config.TextColumn("확인내역", disabled=not is_master),
            "불량원인": st.column_config.TextColumn("불량원인", disabled=not is_master),
            "수리내역": st.column_config.TextColumn("수리내역", disabled=not is_master),
            "완료일자": st.column_config.TextColumn("완료일자", disabled=not is_master),
            "인계일자": st.column_config.TextColumn("인계일자", disabled=not is_master),
            "비고": st.column_config.TextColumn("비고", disabled=not is_master)
        }

        editor_df = view_df[['선택'] + EXCEL_FIELDS].copy()
        
        edited_df = st.data_editor(
            editor_df,
            column_config=column_config,
            hide_index=True,
            height=850,
            num_rows="dynamic" if is_master else "fixed",
            key="stable_as_table_editor_v6"
        )

        for idx in edited_df.index:
            st.session_state["display_df"].loc[idx, ['선택'] + EXCEL_FIELDS] = edited_df.loc[idx, ['선택'] + EXCEL_FIELDS].values

        b_col1, b_col2, b_col3 = st.columns([3, 3, 4])
        
        with b_col1:
            if is_master:
                if st.button("💾 표에서 수정한 내용 DB에 영구 저장"):
                    try:
                        existing_db_df = load_fresh_db_data(current_order)
                        photo_map = {}
                        if not existing_db_df.empty:
                            for _, row in existing_db_df.iterrows():
                                no_key = str(row.get('NO.', '')).strip()
                                sn_key = str(row.get('제품 S/N', '')).strip()
                                photo_map[(no_key, sn_key)] = (row.get('확인내역_사진', ''), row.get('수리내역_사진', ''))

                        raw_save_df = st.session_state["display_df"][EXCEL_FIELDS].copy()
                        if current_order.startswith("최신순"):
                            raw_save_df = raw_save_df.iloc[::-1].reset_index(drop=True)

                        meaningful_cols = ['접수일', '프로젝트', '제품명', '제품 S/N', '접수내역', '확인내역']
                        valid_rows = []
                        today_str = datetime.now().strftime('%Y-%m-%d')

                        for _, row in raw_save_df.iterrows():
                            row_dict = row.to_dict()
                            has_data = any(str(row_dict.get(col, '')).strip() != '' and str(row_dict.get(col, '')).strip().lower() not in ['nan', 'none'] for col in meaningful_cols)

                            if has_data:
                                no_key = str(row_dict.get('NO.', '')).strip()
                                sn_key = str(row_dict.get('제품 S/N', '')).strip()
                                
                                if (no_key, sn_key) in photo_map:
                                    if not str(row_dict.get('확인내역_사진', '')).strip():
                                        row_dict['확인내역_사진'] = photo_map[(no_key, sn_key)][0]
                                    if not str(row_dict.get('수리내역_사진', '')).strip():
                                        row_dict['수리내역_사진'] = photo_map[(no_key, sn_key)][1]

                                for test_k, date_k in INSPECT_1ST_PAIRS.items():
                                    cur_val = str(row_dict.get(test_k, '')).strip().upper()
                                    cur_date = clean_date_str(row_dict.get(date_k, ''))
                                    if cur_val in ['PASS', 'FAIL']:
                                        if not cur_date: row_dict[date_k] = today_str
                                    else:
                                        if not cur_val: row_dict[date_k] = ""

                                for test_k, date_k in INSPECT_2ND_PAIRS.items():
                                    cur_val = str(row_dict.get(test_k, '')).strip().upper()
                                    cur_date = clean_date_str(row_dict.get(date_k, ''))
                                    if cur_val in ['PASS', 'FAIL']:
                                        if not cur_date: row_dict[date_k] = today_str
                                    else:
                                        if not cur_val: row_dict[date_k] = ""

                                valid_rows.append(row_dict)

                        if not valid_rows:
                            st.warning("저장할 유효 데이터가 없습니다.")
                        else:
                            processed_df = pd.DataFrame(valid_rows, columns=EXCEL_FIELDS)

                            int_nos = [safe_int_no(v) for v in processed_df['NO.']]
                            valid_existing = [x for x in int_nos if x is not None]
                            max_existing_no = max(valid_existing) if valid_existing else 0

                            final_no_list = []
                            for v in int_nos:
                                if v is None:
                                    max_existing_no += 1
                                    final_no_list.append(max_existing_no)
                                else:
                                    final_no_list.append(v)

                            processed_df['final_no'] = final_no_list
                            processed_df['NO.'] = [str(x) for x in final_no_list]
                            processed_df.sort_values(by='final_no', ascending=True, inplace=True)
                            processed_df.drop(columns=['final_no'], inplace=True)

                            for col in DATE_FIELDS:
                                if col in processed_df.columns:
                                    processed_df[col] = processed_df[col].apply(clean_date_str)

                            processed_df = calculate_reception_counts(processed_df)
                            final_save_df = processed_df[EXCEL_FIELDS].fillna("").astype(str)
                            final_save_df.to_sql("as_data", conn, if_exists="replace", index=False)

                            st.session_state["display_df"] = load_fresh_db_data(current_order)
                            if '선택' not in st.session_state["display_df"].columns:
                                st.session_state["display_df"].insert(0, '선택', False)
                                
                            st.toast("✅ 저장 완료! 데이터가 데이터베이스에 영구 저장되었습니다.", icon="💾")
                            st.rerun()

                    except Exception as save_err:
                        st.error(f"저장 중 오류 발생: {save_err}")
            else:
                st.caption("🔒 일반 사용자는 표 수정 및 DB 저장이 제한됩니다.")

        with b_col2:
            if is_master:
                if st.button("🗑️ 선택한 행 DB에서 완전 삭제"):
                    delete_targets = edited_df[edited_df['선택'] == True]
                    if delete_targets.empty:
                        st.warning("삭제할 행을 앞쪽 [선택] 체크박스로 먼저 지정해 주세요.")
                    else:
                        try:
                            sel_indices = delete_targets.index.tolist()
                            remaining_df = st.session_state["display_df"].drop(index=sel_indices).copy()
                            clean_rem_df = remaining_df[EXCEL_FIELDS].copy()

                            meaningful_cols = ['접수일', '프로젝트', '제품명', '제품 S/N', '접수내역', '확인내역']
                            valid_remaining = []
                            for _, row in clean_rem_df.iterrows():
                                if any(str(row.get(c, '')).strip() not in ['', 'nan', 'None'] for c in meaningful_cols):
                                    valid_remaining.append(row)

                            if valid_remaining:
                                final_db_df = pd.DataFrame(valid_remaining, columns=EXCEL_FIELDS).fillna("").astype(str)
                            else:
                                final_db_df = pd.DataFrame(columns=EXCEL_FIELDS)

                            final_db_df.to_sql("as_data", conn, if_exists="replace", index=False)
                            
                            st.session_state["display_df"] = load_fresh_db_data(current_order)
                            if '선택' not in st.session_state["display_df"].columns:
                                st.session_state["display_df"].insert(0, '선택', False)
                                
                            st.toast(f"🗑️ {len(delete_targets)}건 삭제 완료", icon="✅")
                            st.rerun()
                        except Exception as del_err:
                            st.error(f"삭제 처리 중 오류 발생: {del_err}")
            else:
                st.caption("🔒 일반 사용자는 행 삭제가 제한됩니다.")

        with b_col3:
            fresh_ledger = load_fresh_db_data(current_order)[EXCEL_FIELDS]
            ledger_excel_bytes = make_excel_bytes_with_merge(fresh_ledger)
            st.download_button(
                label="📥 AS관리대장 엑셀 다운로드",
                data=ledger_excel_bytes,
                file_name=f"AS관리대장_{datetime.now().strftime('%y%m%d_%H%M%S')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        selected_rows_df = edited_df[edited_df['선택'] == True]
        
        st.markdown("---")
        if not selected_rows_df.empty:
            selected_rows = [selected_rows_df.iloc[i].to_dict() for i in range(len(selected_rows_df))]
            
            st.subheader(f"📑 수리 REPORT 엑셀 다운로드 (총 {len(selected_rows)}건 선택됨)")
            summary_labels = [f"NO.{r.get('NO.', '')} ({r.get('제품명', '')})" for r in selected_rows]
            st.caption(f"선택 항목: {', '.join(summary_labels[:8])}{' 외 ' + str(len(summary_labels)-8) + '건' if len(summary_labels) > 8 else ''}")

            try:
                multi_excel_bytes = generate_multi_repair_excel(selected_rows)
                st.download_button(
                    label=f"📥 선택한 {len(selected_rows)}건 수리 REPORT 엑셀 다운로드",
                    data=multi_excel_bytes,
                    file_name=f"수리Report_일괄_{datetime.now().strftime('%y%m%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary"
                )
            except Exception as xl_err:
                st.error(f"엑셀 생성 중 오류: {xl_err}")
        else:
            st.info("👆 위 표의 맨 앞 [선택] 체크박스를 클릭하면 해당 건의 수리 REPORT 엑셀 다운로드 버튼이 활성화됩니다.")

    # =============================================================
    # [TAB 2: 2026년 KPI 데이터 분석 대시보드]
    # =============================================================
    with tab2:
        st.subheader("📊 2026년 품질 및 AS 운영 KPI 대시보드")

        kpi_base_df = load_fresh_db_data(current_order)
        kpi_base_df['접수일_dt'] = pd.to_datetime(kpi_base_df['접수일'].replace('', pd.NA), errors='coerce')
        valid_dates = kpi_base_df['접수일_dt'].dropna()

        available_years = sorted(valid_dates.dt.year.unique().tolist(), reverse=True) if not valid_dates.empty else [2026]
        if 2026 not in available_years: available_years.insert(0, 2026)

        f_col1, f_col2, f_col3, f_col4, f_col5 = st.columns([2.2, 2.0, 2.0, 2.2, 3.6])

        with f_col1:
            period_mode = st.selectbox("📅 분석 주기", ["연간 분석", "월간 분석", "직접 기간 지정"], index=1)

        sel_year = available_years[0]
        sel_month = datetime.now().month

        if period_mode == "연간 분석":
            with f_col2: sel_year = st.selectbox("연도 선택", available_years, index=0)
            start_date = datetime(sel_year, 1, 1).date()
            end_date = datetime(sel_year, 12, 31).date()
            with f_col3: st.info(f"{sel_year}년 전체")

        elif period_mode == "월간 분석":
            with f_col2: sel_year = st.selectbox("연도 선택", available_years, index=0)
            with f_col3:
                default_m_idx = (datetime.now().month - 1) if sel_year == datetime.now().year else 0
                sel_month = st.selectbox("월 선택", [f"{m}월" for m in range(1, 13)], index=default_m_idx)
                month_num = int(sel_month.replace("월", ""))
            start_date = datetime(sel_year, month_num, 1).date()
            if month_num == 12: end_date = datetime(sel_year, 12, 31).date()
            else: end_date = (datetime(sel_year, month_num + 1, 1) - pd.Timedelta(days=1)).date()

        else:
            min_d = valid_dates.min().date() if not valid_dates.empty else datetime(2026, 1, 1).date()
            max_d = valid_dates.max().date() if not valid_dates.empty else datetime.now().date()
            with f_col2: start_date = st.date_input("시작일", value=min_d)
            with f_col3: end_date = st.date_input("종료일", value=max_d)

        with f_col4:
            cost_options = ["전체 (유/무상)", "무상", "유상", "미기재/기타"]
            selected_cost = st.selectbox("💰 유/무상 구분", cost_options, index=0)

        with f_col5:
            raw_projects = sorted([str(p).strip() for p in kpi_base_df['프로젝트'].unique() if str(p).strip()])
            proj_options = ["전체 프로젝트"] + raw_projects
            selected_proj_filter = st.multiselect("🏗️ 대상 프로젝트", options=proj_options, default=["전체 프로젝트"])

        kpi_df = kpi_base_df.copy()
        kpi_df = kpi_df[(kpi_df['접수일_dt'].dt.date >= start_date) & (kpi_df['접수일_dt'].dt.date <= end_date)]

        if selected_cost == "무상":
            kpi_df = kpi_df[kpi_df['유/무상'].astype(str).str.strip().str.contains("무상", na=False)]
        elif selected_cost == "유상":
            kpi_df = kpi_df[
                kpi_df['유/무상'].astype(str).str.strip().str.contains("유상", na=False) &
                ~kpi_df['유/무상'].astype(str).str.strip().str.contains("무상", na=False)
            ]
        elif selected_cost == "미기재/기타":
            kpi_df = kpi_df[~kpi_df['유/무상'].astype(str).str.strip().isin(["유상", "무상"])]

        if "전체 프로젝트" not in selected_proj_filter and selected_proj_filter:
            kpi_df = kpi_df[kpi_df['프로젝트'].isin(selected_proj_filter)]

        st.caption(f"📌 **현재 분석 기준**: {start_date} ~ {end_date} | 구분: **{selected_cost}** | 총 **{len(kpi_df):,}건**")
        st.markdown("---")

        total_receptions = len(kpi_df)

        valid_tat_df = kpi_df.copy()
        valid_tat_df['완료일자_dt'] = pd.to_datetime(valid_tat_df['완료일자'].replace('', pd.NA), errors='coerce')
        valid_dates_df = valid_tat_df.dropna(subset=['접수일_dt', '완료일자_dt'])
        valid_tat_df['리드타임'] = (valid_dates_df['완료일자_dt'] - valid_tat_df['접수일_dt']).dt.days
        valid_tat_df = valid_tat_df[valid_tat_df['리드타임'] >= 0]
        avg_tat = round(valid_tat_df['리드타임'].mean(), 1) if not valid_tat_df.empty else 0.0

        re_reception_count = len(kpi_df[kpi_df['접수횟수'].astype(str).str.strip().isin(['2', '3', '4', '5', '6', '7', '8', '9'])])
        re_rate = round((re_reception_count / total_receptions * 100), 1) if total_receptions > 0 else 0.0

        completed_count = len(kpi_df[kpi_df['완료일자'].str.strip() != ''])
        completion_rate = round((completed_count / total_receptions * 100), 1) if total_receptions > 0 else 0.0

        m1, m2, m3, m4 = st.columns(4)
        m1.metric("총 접수 건수", f"{total_receptions:,} 건")
        m2.metric("평균 수리 소요일 (TAT)", f"{avg_tat} 일", help="접수일로부터 완료일까지의 소요 일수")
        m3.metric("재접수율 (불량 재발)", f"{re_rate} %", f"{re_reception_count}건 재접수")
        m4.metric("수리 완료율", f"{completion_rate} %", f"{completed_count}/{total_receptions}건 완료")

        st.markdown("---")

        c_col1, c_col2 = st.columns(2)

        with c_col1:
            if period_mode == "연간 분석":
                st.markdown("##### 📈 월별 접수 및 완료 추이")
                kpi_df['기준월'] = kpi_df['접수일_dt'].dt.strftime('%m월')
                monthly_recv = kpi_df.groupby('기준월').size().rename("접수")
                trend_chart_df = pd.DataFrame(monthly_recv)
                st.line_chart(trend_chart_df)
            else:
                st.markdown("##### 📈 일별 접수 추이")
                kpi_df['접수일자_str'] = kpi_df['접수일_dt'].dt.strftime('%m-%d')
                daily_recv = kpi_df.groupby('접수일자_str').size().rename("접수 건수")
                st.line_chart(daily_recv)

        with c_col2:
            st.markdown("##### 🔍 불량 원인 순위 (Top 5)")
            cause_series = kpi_df['불량원인'].replace('', '원인 미기재').value_counts().head(5)
            if not cause_series.empty:
                st.bar_chart(cause_series)
            else:
                st.info("데이터가 없습니다.")

        c_col3, c_col4 = st.columns(2)

        with c_col3:
            st.markdown("##### 🏢 프로젝트별 발생 현황 (Top 5)")
            proj_series = kpi_df['프로젝트'].replace('', '기타').value_counts().head(5)
            if not proj_series.empty:
                st.bar_chart(proj_series)
            else:
                st.info("데이터가 없습니다.")

        with c_col4:
            st.markdown("##### ⚙️ 주요 투입 교체 부품 (Top 5)")
            parts_list = []
            for c in ['투입부품1', '투입부품2', '투입부품3', '투입부품4']:
                parts = kpi_df[c].dropna().astype(str).str.strip()
                parts = parts[~parts.isin(['', 'nan', 'None', '-'])]
                parts_list.extend(parts.tolist())
            
            if parts_list:
                part_counts = pd.Series(parts_list).value_counts().head(5)
                st.bar_chart(part_counts)
            else:
                st.info("기록된 부품 내역이 없습니다.")

except Exception as e:
    st.error(f"데이터 로드 및 처리 중 오류 발생: {e}")